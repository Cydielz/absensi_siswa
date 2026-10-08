"""Logika bisnis: validasi data, QR bertanda tangan (HMAC), pencatatan absensi,
antrean WhatsApp (Fonnte) di latar belakang, backup, dan import siswa."""
import csv
import glob
import hmac
import hashlib
import io
import logging
import os
import re
import sqlite3
import threading
from datetime import datetime, timedelta

import requests

import database as dbm

log = logging.getLogger('absensi')

QR_PREFIX = 'ABSEN'
_QR_RE = re.compile(r'^ABSEN:(\d{10}):([0-9a-f]{12})$')
_LEGACY_RE = re.compile(r'^(?:NISN:)?\s*(\d{10})$', re.IGNORECASE)


# ----------------------------------------------------------------- waktu
def now_local(tz):
    return datetime.now(tz)


def now_str(tz):
    return now_local(tz).strftime('%Y-%m-%d %H:%M:%S')


# ------------------------------------------------------------ normalisasi
def normalize_nisn(value):
    """Ambil angka saja; lengkapi nol di depan (Excel sering menghilangkannya).
    Kembalikan string 10 digit atau None bila tidak valid."""
    digits = re.sub(r'\D', '', str(value or ''))
    if not digits or len(digits) > 10:
        return None
    return digits.zfill(10)


def normalize_wa(value):
    """Ubah ke format 62xxxxxxxxxx. '' bila kosong, None bila tidak valid."""
    raw = str(value or '').strip()
    if not raw:
        return ''
    d = re.sub(r'\D', '', raw)
    if d.startswith('62'):
        pass
    elif d.startswith('0'):
        d = '62' + d[1:]
    elif d.startswith('8'):
        d = '62' + d
    else:
        return None
    if not 10 <= len(d) <= 15:
        return None
    return d


# --------------------------------------------------------------------- QR
def qr_signature(secret, nisn, version):
    msg = f'{nisn}:{version}'.encode()
    return hmac.new(secret.encode(), msg, hashlib.sha256).hexdigest()[:12]


def make_qr_payload(secret, nisn, version):
    return f'{QR_PREFIX}:{nisn}:{qr_signature(secret, nisn, version)}'


def verify_qr(conn, cfg, raw):
    """Kembalikan (siswa_row, None) bila valid, atau (None, kode_error)."""
    text = (raw or '').strip()
    m = _QR_RE.match(text)
    if m:
        nisn, sig = m.group(1), m.group(2)
        siswa = conn.execute('SELECT * FROM siswa WHERE nisn=?', (nisn,)).fetchone()
        if siswa is None or not siswa['aktif']:
            return None, 'STUDENT_NOT_FOUND'
        expected = qr_signature(cfg['QR_SECRET'], nisn, siswa['qr_version'])
        if not hmac.compare_digest(sig, expected):
            return None, 'QR_INVALID'
        return siswa, None

    legacy = _LEGACY_RE.match(text)
    if legacy:
        if not cfg.get('ALLOW_LEGACY_NISN'):
            return None, 'QR_LEGACY'
        siswa = conn.execute('SELECT * FROM siswa WHERE nisn=? AND aktif=1', (legacy.group(1),)).fetchone()
        if siswa is None:
            return None, 'STUDENT_NOT_FOUND'
        return siswa, None
    return None, 'QR_INVALID'


# ------------------------------------------------------------- pesan WA
def render_message(template, values):
    """Ganti {placeholder}. Aman: tidak memakai str.format."""
    return re.sub(r'\{(\w+)\}', lambda m: str(values.get(m.group(1), m.group(0))), template)


def enqueue_whatsapp(conn, cfg, siswa, absensi_id, status, jam, tanggal):
    """Masukkan notifikasi ke antrean. Mengembalikan id log, atau None bila tidak dikirim."""
    if dbm.get_setting(conn, 'wa_aktif', '1') != '1':
        return None
    target = (siswa['whatsapp_ortu'] or '').strip()
    if not target:
        return None
    d = datetime.strptime(tanggal, '%Y-%m-%d').strftime('%d-%m-%Y')
    message = render_message(dbm.get_setting(conn, 'wa_template'), {
        'nama': siswa['nama'], 'nisn': siswa['nisn'], 'kelas': siswa['kelas'],
        'status': status.upper(), 'tanggal': d, 'jam': jam,
        'sekolah': dbm.get_setting(conn, 'nama_sekolah'),
    })
    stamp = now_str(cfg['TZ'])
    cur = conn.execute(
        '''INSERT INTO wa_log (siswa_id, absensi_id, target, message, status, attempts,
                               next_attempt_at, created_at, updated_at)
           VALUES (?, ?, ?, ?, 'pending', 0, ?, ?, ?)''',
        (siswa['id'], absensi_id, target, message, stamp, stamp, stamp),
    )
    return cur.lastrowid


def send_fonnte(token, target, message, country_code='62', timeout=20):
    """Kirim satu pesan. Mengembalikan (berhasil, detail).
    Catatan: Fonnte membalas HTTP 200 meski gagal, jadi field `status` di body wajib dicek."""
    try:
        resp = requests.post(
            'https://api.fonnte.com/send',
            headers={'Authorization': token},
            data={'target': target, 'message': message, 'countryCode': country_code},
            timeout=timeout,
        )
        try:
            body = resp.json()
        except ValueError:
            body = {'raw': resp.text[:300]}
        ok = bool(resp.ok and isinstance(body, dict) and body.get('status') is True)
        return ok, str(body)[:500]
    except requests.RequestException as exc:
        return False, f'Jaringan: {exc}'[:500]


def process_wa_queue(cfg, sender=None, limit=20):
    """Proses pesan 'pending'/'retry' yang sudah jatuh tempo. Mengembalikan jumlah diproses."""
    sender = sender or send_fonnte
    tz = cfg['TZ']
    conn = dbm.connect(cfg['DB_PATH'])
    try:
        rows = conn.execute(
            '''SELECT * FROM wa_log WHERE status IN ('pending','retry') AND next_attempt_at <= ?
               ORDER BY id LIMIT ?''', (now_str(tz), limit)).fetchall()
        for r in rows:
            stamp = now_str(tz)
            if not cfg['FONNTE_TOKEN']:
                conn.execute("UPDATE wa_log SET status='failed', last_response=?, updated_at=? WHERE id=?",
                             ('FONNTE_TOKEN belum diatur di .env', stamp, r['id']))
                continue
            ok, detail = sender(cfg['FONNTE_TOKEN'], r['target'], r['message'], cfg['FONNTE_COUNTRY_CODE'])
            attempts = r['attempts'] + 1
            if ok:
                conn.execute("UPDATE wa_log SET status='sent', attempts=?, last_response=?, updated_at=?, "
                             "next_attempt_at=NULL WHERE id=?", (attempts, detail, stamp, r['id']))
            elif attempts >= cfg['WA_MAX_ATTEMPTS']:
                conn.execute("UPDATE wa_log SET status='failed', attempts=?, last_response=?, updated_at=? WHERE id=?",
                             (attempts, detail, stamp, r['id']))
            else:
                nxt = (now_local(tz) + timedelta(seconds=60 * attempts ** 2)).strftime('%Y-%m-%d %H:%M:%S')
                conn.execute("UPDATE wa_log SET status='retry', attempts=?, last_response=?, updated_at=?, "
                             "next_attempt_at=? WHERE id=?", (attempts, detail, stamp, nxt, r['id']))
            conn.commit()
        conn.commit()
        return len(rows)
    finally:
        conn.close()


# ----------------------------------------------------------------- backup
def backup_db(src, dest):
    s = dbm.connect(src)
    d = sqlite3.connect(dest)
    try:
        s.backup(d)
    finally:
        d.close()
        s.close()


def maybe_daily_backup(cfg):
    os.makedirs(cfg['BACKUP_DIR'], exist_ok=True)
    name = f"sekolah-{now_local(cfg['TZ']).strftime('%Y%m%d')}.db"
    dest = os.path.join(cfg['BACKUP_DIR'], name)
    if os.path.exists(dest):
        return None
    backup_db(cfg['DB_PATH'], dest)
    files = sorted(glob.glob(os.path.join(cfg['BACKUP_DIR'], 'sekolah-*.db')))
    for old in files[:-cfg['BACKUP_KEEP']]:
        try:
            os.remove(old)
        except OSError:
            pass
    log.info('Backup harian dibuat: %s', dest)
    return dest


class BackgroundWorker(threading.Thread):
    """Mengirim WA dari antrean dan membuat backup harian, tanpa menghambat /scan."""

    def __init__(self, cfg):
        super().__init__(daemon=True, name='absensi-worker')
        self.cfg = cfg
        self.wake = threading.Event()
        self.stop_event = threading.Event()

    def run(self):
        while not self.stop_event.is_set():
            try:
                process_wa_queue(self.cfg)
                maybe_daily_backup(self.cfg)
            except Exception:  # jangan sampai worker mati
                log.exception('Worker error')
            self.wake.wait(5)
            self.wake.clear()


# ---------------------------------------------------------------- absensi
def student_payload(siswa):
    return {
        'nisn': siswa['nisn'], 'nama': siswa['nama'], 'kelas': siswa['kelas'],
        'foto_url': f"/foto/{siswa['id']}?v={siswa['foto']}" if siswa['foto'] else '',
    }


def record_scan(conn, cfg, siswa, device_id):
    """Catat absensi dari scan. Mengembalikan (http_status, payload)."""
    current = now_local(cfg['TZ'])
    tanggal = current.strftime('%Y-%m-%d')
    jam = current.strftime('%H:%M:%S')
    batas = dbm.get_setting(conn, 'batas_terlambat', '07:15')
    status = 'Terlambat' if jam[:5] > batas else 'Hadir'
    stamp = current.strftime('%Y-%m-%d %H:%M:%S')
    st = student_payload(siswa)

    existing = conn.execute('SELECT * FROM absensi WHERE siswa_id=? AND tanggal=?',
                            (siswa['id'], tanggal)).fetchone()
    if existing and existing['status'] in dbm.STATUS_HADIR:
        return 200, {
            'success': True, 'code': 'ALREADY_ATTENDED',
            'message': f"{siswa['nama']} sudah absen hari ini pada {existing['jam_masuk']}.",
            'student': st,
            'attendance': {'tanggal': tanggal, 'jam_masuk': existing['jam_masuk'], 'status': existing['status']},
        }

    if existing:  # sebelumnya Izin/Sakit/Alpa tetapi ternyata datang
        conn.execute('UPDATE absensi SET status=?, jam_masuk=?, device_id=?, updated_at=? WHERE id=?',
                     (status, jam, device_id, stamp, existing['id']))
        attendance_id = existing['id']
    else:
        cur = conn.execute(
            '''INSERT INTO absensi (siswa_id, tanggal, jam_masuk, status, device_id, created_at, updated_at)
               VALUES (?, ?, ?, ?, ?, ?, ?)''', (siswa['id'], tanggal, jam, status, device_id, stamp, stamp))
        attendance_id = cur.lastrowid

    log_id = enqueue_whatsapp(conn, cfg, siswa, attendance_id, status, jam, tanggal)
    conn.commit()
    return 201, {
        'success': True, 'code': 'ATTENDANCE_RECORDED',
        'message': f"Absensi {siswa['nama']} berhasil dicatat pada {jam}.",
        'student': st,
        'attendance': {'id': attendance_id, 'tanggal': tanggal, 'jam_masuk': jam, 'status': status},
        'whatsapp': {'queued': log_id is not None},
    }


def set_attendance(conn, cfg, siswa_id, tanggal, status, keterangan='', device='admin'):
    """Atur/ubah status absensi manual (izin, sakit, alpa, atau koreksi)."""
    tz = cfg['TZ']
    stamp = now_str(tz)
    today = now_local(tz).strftime('%Y-%m-%d')
    existing = conn.execute('SELECT * FROM absensi WHERE siswa_id=? AND tanggal=?', (siswa_id, tanggal)).fetchone()
    if status in dbm.STATUS_HADIR:
        if existing and existing['jam_masuk'] != '-':
            jam = existing['jam_masuk']
        elif tanggal == today:
            jam = now_local(tz).strftime('%H:%M:%S')
        else:
            jam = '-'
    else:
        jam = '-'
    if existing:
        conn.execute('UPDATE absensi SET status=?, jam_masuk=?, keterangan=?, updated_at=? WHERE id=?',
                     (status, jam, keterangan, stamp, existing['id']))
    else:
        conn.execute(
            '''INSERT INTO absensi (siswa_id, tanggal, jam_masuk, status, device_id, keterangan, created_at, updated_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)''', (siswa_id, tanggal, jam, status, device, keterangan, stamp, stamp))


def mark_alpa(conn, cfg, tanggal, kelas=''):
    """Tandai Alpa untuk siswa aktif yang belum punya catatan pada tanggal tsb."""
    sql = '''SELECT s.id FROM siswa s LEFT JOIN absensi a ON a.siswa_id=s.id AND a.tanggal=?
             WHERE s.aktif=1 AND a.id IS NULL'''
    params = [tanggal]
    if kelas:
        sql += ' AND s.kelas=?'
        params.append(kelas)
    ids = [r['id'] for r in conn.execute(sql, params)]
    stamp = now_str(cfg['TZ'])
    for sid in ids:
        conn.execute(
            '''INSERT INTO absensi (siswa_id, tanggal, jam_masuk, status, device_id, created_at, updated_at)
               VALUES (?, ?, '-', 'Alpa', 'sistem', ?, ?)''', (sid, tanggal, stamp, stamp))
    return len(ids)


# ----------------------------------------------------------------- import
def _cell(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_table(filename, data):
    """Baca .xlsx atau .csv menjadi list baris (list of str)."""
    name = (filename or '').lower()
    if name.endswith('.xlsx'):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        return [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
    if name.endswith('.csv') or name.endswith('.txt'):
        text = data.decode('utf-8-sig', errors='replace')
        first = text.splitlines()[0] if text.strip() else ''
        delim = max(',;\t', key=first.count) if first else ','
        return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), delimiter=delim)]
    raise ValueError('Format file harus .xlsx atau .csv')


def parse_students(rows):
    """Deteksi header (jika ada) dan kembalikan list (nomor_baris, nisn, nama, kelas, wa)."""
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return []
    header = [c.lower() for c in rows[0]]

    def find(*keys):
        for i, h in enumerate(header):
            if any(k in h for k in keys):
                return i
        return None

    idx = {'nisn': find('nisn'), 'nama': find('nama'), 'kelas': find('kelas'),
           'wa': find('wa', 'whatsapp', 'hp', 'telp', 'ortu')}
    has_header = None not in (idx['nisn'], idx['nama'], idx['kelas'])
    if has_header:
        start = 1
    else:
        idx = {'nisn': 0, 'nama': 1, 'kelas': 2, 'wa': 3}
        start = 0
    out = []
    for n, r in enumerate(rows[start:], start=start + 1):
        def g(k):
            i = idx[k]
            return r[i] if i is not None and i < len(r) else ''
        out.append((n, g('nisn'), g('nama'), g('kelas'), g('wa')))
    return out


def import_students(conn, cfg, entries, update_existing=True):
    stamp = now_str(cfg['TZ'])
    added = updated = 0
    errors = []
    for n, nisn_raw, nama, kelas, wa_raw in entries:
        nisn = normalize_nisn(nisn_raw)
        wa = normalize_wa(wa_raw)
        nama, kelas = nama.strip(), kelas.strip()
        if not nisn:
            errors.append(f'Baris {n}: NISN "{nisn_raw}" tidak valid (harus 10 digit angka).')
        elif not nama or not kelas:
            errors.append(f'Baris {n}: nama dan kelas wajib diisi.')
        elif wa is None:
            errors.append(f'Baris {n}: nomor WhatsApp "{wa_raw}" tidak valid.')
        else:
            row = conn.execute('SELECT id FROM siswa WHERE nisn=?', (nisn,)).fetchone()
            if row is None:
                conn.execute('INSERT INTO siswa (nisn,nama,kelas,whatsapp_ortu,created_at) VALUES (?,?,?,?,?)',
                             (nisn, nama, kelas, wa, stamp))
                added += 1
            elif update_existing:
                conn.execute("UPDATE siswa SET nama=?, kelas=?, "
                             "whatsapp_ortu=CASE WHEN ?='' THEN whatsapp_ortu ELSE ? END WHERE id=?",
                             (nama, kelas, wa, wa, row['id']))
                updated += 1
    conn.commit()
    return added, updated, errors
