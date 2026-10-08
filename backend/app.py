"""Absensi Siswa v2 — Flask + SQLite + Fonnte.

Jalankan:  python app.py
"""
import hmac
import io
import json
import logging
import os
import re
import secrets
import socket
import sqlite3
<<<<<<< HEAD
from datetime import datetime
from zoneinfo import ZoneInfo
from flask import Flask, jsonify, request, render_template, redirect, url_for, flash, Response
import requests
from dotenv import load_dotenv
=======
import time
from csv import writer as csv_writer
from datetime import datetime, timedelta
from functools import wraps
from logging.handlers import RotatingFileHandler
from urllib.parse import urlparse
>>>>>>> 200af4a (Absensi Siswa v2)

from flask import (Flask, Response, abort, flash, g, jsonify, redirect, render_template,
                   request, send_file, send_from_directory, session, url_for)
from werkzeug.security import check_password_hash, generate_password_hash

import config as cfgmod
import database as dbm
import services as svc

log = logging.getLogger('absensi')
PER_PAGE = 25
CSRF_EXEMPT = {'scan', 'static'}
NAMA_BULAN = ['', 'Januari', 'Februari', 'Maret', 'April', 'Mei', 'Juni', 'Juli',
              'Agustus', 'September', 'Oktober', 'November', 'Desember']


def setup_logging(cfg):
    log.setLevel(logging.INFO)
    log.propagate = False
    if log.handlers:
        return
    fmt = logging.Formatter('%(asctime)s %(levelname)s %(message)s')
    console = logging.StreamHandler()
    console.setFormatter(fmt)
    log.addHandler(console)
    if cfg['LOG_TO_FILE']:
        os.makedirs(cfg['LOG_DIR'], exist_ok=True)
        fh = RotatingFileHandler(os.path.join(cfg['LOG_DIR'], 'app.log'), maxBytes=1_000_000,
                                 backupCount=5, encoding='utf-8')
        fh.setFormatter(fmt)
        log.addHandler(fh)


def parse_date(value, default):
    try:
        return datetime.strptime(value or '', '%Y-%m-%d').strftime('%Y-%m-%d')
    except ValueError:
        return default


def parse_month(value, default):
    try:
        return datetime.strptime(value or '', '%Y-%m').strftime('%Y-%m')
    except ValueError:
        return default


def safe_cell(v):
    """Cegah CSV/Excel formula injection."""
    s = '' if v is None else str(v)
    return "'" + s if s[:1] in ('=', '+', '-', '@', '\t', '\r') else s


def safe_next(target):
    if not target:
        return None
    p = urlparse(target)
    if p.netloc or p.scheme or not target.startswith('/') or target.startswith('//'):
        return None
    return target


def guess_lan_ip():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(('10.255.255.255', 1))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except OSError:
        return '127.0.0.1'


def create_app(overrides=None):
    cfg = cfgmod.load_config(overrides)
    setup_logging(cfg)
    if not cfg['FLASK_SECRET_KEY']:
        raise RuntimeError('FLASK_SECRET_KEY kosong. Jalankan lewat "python app.py" agar dibuat otomatis.')
    cfg['TZ'] = cfgmod.get_tz(cfg['TIMEZONE'])
    os.makedirs(os.path.join(cfg['UPLOAD_DIR'], 'foto'), exist_ok=True)
    dbm.init_db(cfg['DB_PATH'])

    app = Flask(__name__)
    app.config.update(
        SECRET_KEY=cfg['FLASK_SECRET_KEY'],
        MAX_CONTENT_LENGTH=8 * 1024 * 1024,
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE='Lax',
        SESSION_COOKIE_SECURE=cfg['SESSION_COOKIE_SECURE'],
        PERMANENT_SESSION_LIFETIME=timedelta(hours=cfg['SESSION_HOURS']),
    )
    app.extensions['absensi_cfg'] = cfg
    failed_logins = {}

    # ------------------------------------------------------------ db & util
    def get_db():
        if 'db' not in g:
            g.db = dbm.connect(cfg['DB_PATH'])
        return g.db

    @app.teardown_appcontext
    def close_db(_exc):
        conn = g.pop('db', None)
        if conn is not None:
            conn.close()

    def today_str():
        return svc.now_local(cfg['TZ']).strftime('%Y-%m-%d')

    def kelas_list():
        return [r['kelas'] for r in get_db().execute('SELECT DISTINCT kelas FROM siswa ORDER BY kelas')]

    def back_or(default):
        ref = request.referrer
        if ref and urlparse(ref).netloc == request.host:
            return ref
        return default

    def api_key_ok():
        key = (request.headers.get('X-API-Key') or '').encode()
        expected = (cfg['SCAN_API_KEY'] or '').encode()
        return bool(expected) and hmac.compare_digest(key, expected)

    # ----------------------------------------------------------------- CSRF
    def csrf_token():
        if '_csrf' not in session:
            session['_csrf'] = secrets.token_hex(16)
        return session['_csrf']

    @app.before_request
    def csrf_protect():
        if request.method in ('POST', 'PUT', 'PATCH', 'DELETE') and request.endpoint not in CSRF_EXEMPT:
            sent = request.form.get('csrf_token') or request.headers.get('X-CSRF-Token') or ''
            expected = session.get('_csrf', '')
            if not sent or not expected or not hmac.compare_digest(sent.encode(), expected.encode()):
                abort(400, 'Sesi formulir kedaluwarsa. Muat ulang halaman lalu coba lagi.')

    @app.context_processor
    def inject_globals():
        name = 'Absensi Siswa'
        try:
            name = dbm.get_setting(get_db(), 'nama_sekolah', name)
        except Exception:
            pass
        return {'csrf_token': csrf_token, 'nama_sekolah': name, 'is_admin': bool(session.get('admin')),
                'bulan_nama': NAMA_BULAN}

    @app.after_request
    def security_headers(resp):
        resp.headers.setdefault('X-Content-Type-Options', 'nosniff')
        resp.headers.setdefault('X-Frame-Options', 'DENY')
        resp.headers.setdefault('Referrer-Policy', 'same-origin')
        return resp

    def login_required(f):
        @wraps(f)
        def wrapper(*a, **kw):
            if not session.get('admin'):
                if request.path.startswith('/api/'):
                    return jsonify({'success': False, 'code': 'UNAUTHORIZED', 'message': 'Perlu login.'}), 401
                nxt = request.full_path.rstrip('?') if request.method == 'GET' else None
                return redirect(url_for('login', next=nxt))
            return f(*a, **kw)
        return wrapper

    @app.errorhandler(400)
    @app.errorhandler(404)
    @app.errorhandler(413)
    def handle_errors(err):
        code = getattr(err, 'code', 500)
        msg = {413: 'File terlalu besar (maksimal 8 MB).', 404: 'Halaman tidak ditemukan.'}.get(
            code, getattr(err, 'description', 'Permintaan tidak valid.'))
        if request.path.startswith(('/scan', '/api/')):
            return jsonify({'success': False, 'code': f'HTTP_{code}', 'message': msg}), code
        return render_template('error.html', code=code, message=msg), code

    # ----------------------------------------------------------------- auth
    def check_login(username, password):
        user_ok = hmac.compare_digest(username.encode(), cfg['ADMIN_USERNAME'].encode())
        stored = dbm.get_setting(get_db(), 'admin_password_hash', '')
        if stored:
            pw_ok = check_password_hash(stored, password)
        else:
            expected = cfg['ADMIN_PASSWORD']
            pw_ok = bool(expected) and hmac.compare_digest(password.encode(), expected.encode())
        return user_ok and pw_ok

    @app.route('/login', methods=['GET', 'POST'])
    def login():
        if request.method == 'POST':
            ip = request.remote_addr or '-'
            attempts = [t for t in failed_logins.get(ip, []) if time.monotonic() - t < 300]
            failed_logins[ip] = attempts
            if len(attempts) >= 5:
                flash('Terlalu banyak percobaan gagal. Coba lagi dalam 5 menit.', 'error')
                return render_template('login.html'), 429
            if check_login(request.form.get('username', ''), request.form.get('password', '')):
                failed_logins.pop(ip, None)
                session.clear()
                session['admin'] = True
                session.permanent = True
                log.info('Login admin berhasil dari %s', ip)
                return redirect(safe_next(request.args.get('next')) or url_for('dashboard'))
            attempts.append(time.monotonic())
            log.warning('Login admin gagal dari %s', ip)
            flash('Username atau password salah.', 'error')
        return render_template('login.html')

    @app.post('/logout')
    def logout():
        session.clear()
        return redirect(url_for('login'))

    # ------------------------------------------------------- API perangkat
    @app.get('/api/health')
    def health():
        return jsonify({'ok': True})

    @app.get('/api/device/ping')
    def device_ping():
        if not api_key_ok():
            return jsonify({'success': False, 'code': 'UNAUTHORIZED', 'message': 'API key salah.'}), 401
        db = get_db()
        return jsonify({'ok': True, 'sekolah': dbm.get_setting(db, 'nama_sekolah'),
                        'tanggal': today_str(), 'batas_terlambat': dbm.get_setting(db, 'batas_terlambat')})

    @app.post('/scan')
    def scan():
        if not api_key_ok():
            return jsonify({'success': False, 'code': 'UNAUTHORIZED',
                            'message': 'API key tidak valid. Periksa Pengaturan di aplikasi.'}), 401
        data = request.get_json(silent=True) or {}
        raw = str(data.get('qr') or data.get('nisn') or '').strip()[:200]
        device_id = str(data.get('device_id', '')).strip()[:100]
        if not raw:
            return jsonify({'success': False, 'code': 'QR_REQUIRED', 'message': 'Isi QR kosong.'}), 400
        db = get_db()
        siswa, err = svc.verify_qr(db, cfg, raw)
        if err:
            msgs = {
                'STUDENT_NOT_FOUND': ('Siswa tidak terdaftar atau tidak aktif.', 404),
                'QR_INVALID': ('QR tidak valid / bukan kartu resmi sekolah.', 422),
                'QR_LEGACY': ('Kartu QR lama (hanya NISN) tidak lagi diterima. Cetak ulang kartu.', 422),
            }
            message, code = msgs[err]
            log.warning('Scan ditolak (%s) dari perangkat %s', err, device_id)
            return jsonify({'success': False, 'code': err, 'message': message}), code
        try:
            status, payload = svc.record_scan(db, cfg, siswa, device_id)
        except Exception:
            log.exception('Gagal mencatat absensi')
            return jsonify({'success': False, 'code': 'SERVER_ERROR', 'message': 'Kesalahan server.'}), 500
        worker = app.extensions.get('worker')
        if worker:
            worker.wake.set()
        return jsonify(payload), status

    @app.get('/foto/<int:siswa_id>')
    def foto(siswa_id):
        if not (session.get('admin') or api_key_ok()):
            abort(404)
        row = get_db().execute('SELECT foto FROM siswa WHERE id=?', (siswa_id,)).fetchone()
        if not row or not row['foto']:
            abort(404)
        resp = send_from_directory(os.path.join(cfg['UPLOAD_DIR'], 'foto'), row['foto'])
        resp.cache_control.private = True
        resp.cache_control.max_age = 3600
        return resp

    @app.get('/api/attendance/today')
    def attendance_today():
        if not (session.get('admin') or api_key_ok()):
            return jsonify({'success': False, 'code': 'UNAUTHORIZED', 'message': 'Perlu login/API key.'}), 401
        tanggal = today_str()
        rows = get_db().execute('''
            SELECT a.id, a.tanggal, a.jam_masuk, a.status, a.device_id, s.nisn, s.nama, s.kelas
            FROM absensi a JOIN siswa s ON s.id=a.siswa_id WHERE a.tanggal=? ORDER BY a.jam_masuk''',
                                (tanggal,)).fetchall()
        return jsonify({'date': tanggal, 'data': [dict(r) for r in rows]})

    # ------------------------------------------------------------ dashboard
    @app.get('/')
    @login_required
    def dashboard():
        tanggal = parse_date(request.args.get('tanggal'), today_str())
        kelas = request.args.get('kelas', '').strip()
        sql = '''SELECT s.id, s.nisn, s.nama, s.kelas, a.status, a.jam_masuk, a.keterangan
                 FROM siswa s LEFT JOIN absensi a ON a.siswa_id=s.id AND a.tanggal=?
                 WHERE s.aktif=1'''
        params = [tanggal]
        if kelas:
            sql += ' AND s.kelas=?'
            params.append(kelas)
        sql += ' ORDER BY s.kelas, s.nama COLLATE NOCASE'
        rows = get_db().execute(sql, params).fetchall()
        counts = {k: 0 for k in dbm.STATUS_SEMUA}
        counts['Belum'] = 0
        for r in rows:
            counts[r['status'] or 'Belum'] += 1
        return render_template('dashboard.html', rows=rows, counts=counts, tanggal=tanggal, kelas=kelas,
                               kelas_options=kelas_list(), is_today=(tanggal == today_str()),
                               statuses=dbm.STATUS_SEMUA, total=len(rows))

    @app.post('/admin/absensi/set')
    @login_required
    def absensi_set():
        db = get_db()
        tanggal = parse_date(request.form.get('tanggal'), '')
        status = request.form.get('status', '')
        back = {'tanggal': tanggal or today_str(), 'kelas': request.form.get('back_kelas', '')}
        try:
            siswa_id = int(request.form.get('siswa_id', ''))
        except ValueError:
            siswa_id = 0
        if (not tanggal or tanggal > today_str() or status not in dbm.STATUS_SEMUA
                or not db.execute('SELECT 1 FROM siswa WHERE id=?', (siswa_id,)).fetchone()):
            flash('Data tidak valid: pilih status dan pastikan tanggal tidak di masa depan.', 'error')
        else:
            svc.set_attendance(db, cfg, siswa_id, tanggal, status, request.form.get('keterangan', '').strip()[:200])
            db.commit()
            flash('Status absensi disimpan.', 'success')
        return redirect(url_for('dashboard', **back))

    @app.post('/admin/absensi/tandai-alpa')
    @login_required
    def tandai_alpa():
        tanggal = parse_date(request.form.get('tanggal'), '')
        kelas = request.form.get('kelas', '').strip()
        if not tanggal or tanggal > today_str():
            flash('Tanggal tidak valid.', 'error')
            return redirect(url_for('dashboard'))
        db = get_db()
        n = svc.mark_alpa(db, cfg, tanggal, kelas)
        db.commit()
        flash(f'{n} siswa yang belum absen ditandai Alpa.', 'success')
        return redirect(url_for('dashboard', tanggal=tanggal, kelas=kelas))

    # ---------------------------------------------------------------- siswa
    def save_photo(siswa_id, storage):
        from PIL import Image, ImageOps
        Image.MAX_IMAGE_PIXELS = 40_000_000
        try:
            img = Image.open(io.BytesIO(storage.read()))
            img = ImageOps.exif_transpose(img).convert('RGB')
            img.thumbnail((400, 400))
            name = f'{siswa_id}.jpg'
            img.save(os.path.join(cfg['UPLOAD_DIR'], 'foto', name), 'JPEG', quality=85)
            return name
        except Exception as exc:
            raise ValueError('File foto bukan gambar yang valid (gunakan JPG/PNG).') from exc

    def remove_photo(siswa_id):
        try:
            os.remove(os.path.join(cfg['UPLOAD_DIR'], 'foto', f'{siswa_id}.jpg'))
        except OSError:
            pass

    @app.get('/admin/siswa')
    @login_required
    def siswa_list():
        q = request.args.get('q', '').strip()
        kelas = request.args.get('kelas', '').strip()
        aktif = request.args.get('aktif', '')
        where, params = [], []
        if q:
            where.append('(nama LIKE ? OR nisn LIKE ?)')
            params += [f'%{q}%', f'%{q}%']
        if kelas:
            where.append('kelas=?')
            params.append(kelas)
        if aktif in ('0', '1'):
            where.append('aktif=?')
            params.append(int(aktif))
        clause = ('WHERE ' + ' AND '.join(where)) if where else ''
        db = get_db()
        total = db.execute(f'SELECT COUNT(*) c FROM siswa {clause}', params).fetchone()['c']
        pages = max(1, -(-total // PER_PAGE))
        try:
            page = min(max(1, int(request.args.get('page', 1))), pages)
        except ValueError:
            page = 1
        rows = db.execute(f'SELECT * FROM siswa {clause} ORDER BY kelas, nama COLLATE NOCASE LIMIT ? OFFSET ?',
                          params + [PER_PAGE, (page - 1) * PER_PAGE]).fetchall()
        return render_template('siswa.html', rows=rows, q=q, kelas=kelas, aktif=aktif, page=page,
                               pages=pages, total=total, kelas_options=kelas_list())

    def siswa_form_values(form):
        nisn = svc.normalize_nisn(form.get('nisn'))
        nama = form.get('nama', '').strip()[:100]
        kelas = form.get('kelas', '').strip()[:30]
        wa = svc.normalize_wa(form.get('whatsapp_ortu'))
        if not nisn:
            return None, 'NISN harus 10 digit angka.'
        if not nama or not kelas:
            return None, 'Nama dan kelas wajib diisi.'
        if wa is None:
            return None, 'Nomor WhatsApp tidak valid (contoh: 081234567890).'
        return (nisn, nama, kelas, wa), None

    @app.post('/admin/siswa')
    @login_required
    def siswa_add():
        vals, err = siswa_form_values(request.form)
        if err:
            flash(err, 'error')
            return redirect(url_for('siswa_list'))
        db = get_db()
        try:
            db.execute('INSERT INTO siswa (nisn,nama,kelas,whatsapp_ortu,created_at) VALUES (?,?,?,?,?)',
                       (*vals, svc.now_str(cfg['TZ'])))
            db.commit()
            flash('Siswa berhasil ditambahkan.', 'success')
        except sqlite3.IntegrityError:
            flash('NISN sudah terdaftar.', 'error')
        return redirect(url_for('siswa_list'))

    @app.route('/admin/siswa/<int:siswa_id>/edit', methods=['GET', 'POST'])
    @login_required
    def siswa_edit(siswa_id):
        db = get_db()
        s = db.execute('SELECT * FROM siswa WHERE id=?', (siswa_id,)).fetchone()
        if not s:
            abort(404)
        if request.method == 'POST':
            vals, err = siswa_form_values(request.form)
            if err:
                flash(err, 'error')
                return redirect(url_for('siswa_edit', siswa_id=siswa_id))
            dup = db.execute('SELECT 1 FROM siswa WHERE nisn=? AND id<>?', (vals[0], siswa_id)).fetchone()
            if dup:
                flash('NISN sudah dipakai siswa lain.', 'error')
                return redirect(url_for('siswa_edit', siswa_id=siswa_id))
            foto_name = s['foto']
            try:
                upload = request.files.get('foto')
                if upload and upload.filename:
                    foto_name = save_photo(siswa_id, upload)
                elif request.form.get('hapus_foto'):
                    remove_photo(siswa_id)
                    foto_name = None
            except ValueError as exc:
                flash(str(exc), 'error')
                return redirect(url_for('siswa_edit', siswa_id=siswa_id))
            db.execute('UPDATE siswa SET nisn=?, nama=?, kelas=?, whatsapp_ortu=?, foto=? WHERE id=?',
                       (*vals, foto_name, siswa_id))
            db.commit()
            msg = 'Data siswa diperbarui.'
            if vals[0] != s['nisn']:
                msg += ' NISN berubah: cetak ulang kartu QR siswa ini.'
            flash(msg, 'success')
            return redirect(url_for('siswa_list'))
        return render_template('siswa_edit.html', s=s)

    @app.post('/admin/siswa/<int:siswa_id>/toggle')
    @login_required
    def siswa_toggle(siswa_id):
        db = get_db()
        db.execute('UPDATE siswa SET aktif = 1 - aktif WHERE id=?', (siswa_id,))
        db.commit()
        return redirect(back_or(url_for('siswa_list')))

    @app.post('/admin/siswa/<int:siswa_id>/reset-kartu')
    @login_required
    def siswa_reset_kartu(siswa_id):
        db = get_db()
        db.execute('UPDATE siswa SET qr_version = qr_version + 1 WHERE id=?', (siswa_id,))
        db.commit()
        flash('Kartu lama dinonaktifkan. Cetak kartu baru dari menu Kartu QR.', 'success')
        return redirect(back_or(url_for('siswa_list')))

    @app.post('/admin/siswa/<int:siswa_id>/delete')
    @login_required
    def siswa_delete(siswa_id):
        db = get_db()
        db.execute('DELETE FROM wa_log WHERE siswa_id=?', (siswa_id,))
        db.execute('DELETE FROM siswa WHERE id=?', (siswa_id,))
        db.commit()
        remove_photo(siswa_id)
        flash('Siswa beserta riwayat absensinya dihapus.', 'success')
        return redirect(url_for('siswa_list'))

    @app.post('/admin/siswa/kelas-aksi')
    @login_required
    def kelas_aksi():
        aksi = request.form.get('aksi')
        dari = request.form.get('dari', '').strip()
        ke = request.form.get('ke', '').strip()[:30]
        db = get_db()
        if not dari:
            flash('Pilih kelas asal.', 'error')
        elif aksi == 'pindah' and ke:
            n = db.execute('UPDATE siswa SET kelas=? WHERE kelas=?', (ke, dari)).rowcount
            db.commit()
            flash(f'{n} siswa dipindahkan dari {dari} ke {ke}.', 'success')
        elif aksi == 'nonaktifkan':
            n = db.execute('UPDATE siswa SET aktif=0 WHERE kelas=?', (dari,)).rowcount
            db.commit()
            flash(f'{n} siswa kelas {dari} dinonaktifkan (mis. lulus).', 'success')
        else:
            flash('Isi kelas tujuan.', 'error')
        return redirect(url_for('siswa_list'))

    # --------------------------------------------------------------- import
    def build_xlsx(headers, rows, text_cols=()):
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([safe_cell(v) if isinstance(v, str) else v for v in r])
        for idx in text_cols:
            for cell in ws.iter_rows(min_row=2, min_col=idx + 1, max_col=idx + 1):
                cell[0].number_format = '@'
        for i, h in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(12, len(str(h)) + 4)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def table_response(headers, rows, fmt, base, text_cols=()):
        if fmt == 'csv':
            out = io.StringIO()
            w = csv_writer(out)
            w.writerow(headers)
            for r in rows:
                w.writerow([safe_cell(v) for v in r])
            data = ('\ufeff' + out.getvalue()).encode('utf-8')
            return Response(data, mimetype='text/csv; charset=utf-8',
                            headers={'Content-Disposition': f'attachment; filename="{base}.csv"'})
        return send_file(build_xlsx(headers, rows, text_cols), as_attachment=True, download_name=f'{base}.xlsx',
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    @app.route('/admin/import', methods=['GET', 'POST'])
    @login_required
    def import_page():
        report = None
        if request.method == 'POST':
            f = request.files.get('berkas')
            if not f or not f.filename:
                flash('Pilih file .xlsx atau .csv terlebih dahulu.', 'error')
            else:
                try:
                    entries = svc.parse_students(svc.read_table(f.filename, f.read()))
                    if not entries:
                        flash('File kosong.', 'error')
                    else:
                        db = get_db()
                        added, updated, errors = svc.import_students(
                            db, cfg, entries, update_existing=bool(request.form.get('perbarui')))
                        report = {'added': added, 'updated': updated, 'errors': errors[:50],
                                  'error_count': len(errors)}
                        log.info('Import siswa: +%s ~%s !%s', added, updated, len(errors))
                except ValueError as exc:
                    flash(str(exc), 'error')
                except Exception:
                    log.exception('Gagal import')
                    flash('File tidak dapat dibaca. Gunakan template yang disediakan.', 'error')
        return render_template('import.html', report=report)

    @app.get('/admin/import/template')
    @login_required
    def import_template():
        return table_response(['NISN', 'Nama', 'Kelas', 'WhatsApp Ortu'],
                              [['0012345678', 'Contoh Siswa', 'X IPA 1', '081234567890']],
                              request.args.get('format', 'xlsx'), 'template-siswa', text_cols=(0, 3))

    # ------------------------------------------------------------- kartu QR
    def qr_svg(payload):
        import segno
        return segno.make(payload, error='m').svg_inline(scale=4, border=0, omitsize=True)

    @app.get('/admin/kartu')
    @login_required
    def kartu():
        kelas = request.args.get('kelas', '').strip()
        sid = request.args.get('id', type=int)
        sql, params = 'SELECT * FROM siswa WHERE aktif=1', []
        if sid:
            sql += ' AND id=?'
            params.append(sid)
        elif kelas:
            sql += ' AND kelas=?'
            params.append(kelas)
        sql += ' ORDER BY kelas, nama COLLATE NOCASE'
        cards = [{'s': s, 'svg': qr_svg(svc.make_qr_payload(cfg['QR_SECRET'], s['nisn'], s['qr_version']))}
                 for s in get_db().execute(sql, params)]
        return render_template('kartu.html', cards=cards, kelas=kelas, sid=sid, kelas_options=kelas_list())

    @app.get('/api/qr/<nisn>')
    @login_required
    def qr_payload(nisn):
        s = get_db().execute('SELECT * FROM siswa WHERE nisn=? AND aktif=1', (nisn,)).fetchone()
        if not s:
            return jsonify({'success': False, 'message': 'Siswa tidak ditemukan.'}), 404
        return jsonify({'success': True, 'qr_value': svc.make_qr_payload(cfg['QR_SECRET'], s['nisn'], s['qr_version']),
                        'student': {'nisn': s['nisn'], 'nama': s['nama'], 'kelas': s['kelas']}})

    # ---------------------------------------------------------------- rekap
    def rekap_data(bulan, kelas):
        db = get_db()
        sql = '''SELECT s.id, s.nisn, s.nama, s.kelas,
                   COALESCE(SUM(a.status='Hadir'),0) h, COALESCE(SUM(a.status='Terlambat'),0) t,
                   COALESCE(SUM(a.status='Izin'),0) i, COALESCE(SUM(a.status='Sakit'),0) sk,
                   COALESCE(SUM(a.status='Alpa'),0) al
                 FROM siswa s LEFT JOIN absensi a ON a.siswa_id=s.id AND a.tanggal BETWEEN ? AND ?
                 WHERE s.aktif=1'''
        params = [f'{bulan}-01', f'{bulan}-31']
        if kelas:
            sql += ' AND s.kelas=?'
            params.append(kelas)
        sql += ' GROUP BY s.id ORDER BY s.kelas, s.nama COLLATE NOCASE'
        hari = db.execute('''SELECT COUNT(DISTINCT tanggal) c FROM absensi
                             WHERE tanggal BETWEEN ? AND ? AND status IN ('Hadir','Terlambat')''',
                          (f'{bulan}-01', f'{bulan}-31')).fetchone()['c']
        out = []
        for r in db.execute(sql, params):
            pct = round(min(100.0, (r['h'] + r['t']) * 100 / hari), 1) if hari else 0.0
            out.append({**dict(r), 'persen': pct})
        return out, hari

    @app.get('/admin/rekap')
    @login_required
    def rekap():
        bulan = parse_month(request.args.get('bulan'), today_str()[:7])
        kelas = request.args.get('kelas', '').strip()
        data, hari = rekap_data(bulan, kelas)
        return render_template('rekap.html', data=data, hari=hari, bulan=bulan, kelas=kelas,
                               kelas_options=kelas_list(),
                               judul=f'{NAMA_BULAN[int(bulan[5:])]} {bulan[:4]}')

    @app.get('/admin/rekap/export')
    @login_required
    def rekap_export():
        bulan = parse_month(request.args.get('bulan'), today_str()[:7])
        kelas = request.args.get('kelas', '').strip()
        data, hari = rekap_data(bulan, kelas)
        rows = [[d['nisn'], d['nama'], d['kelas'], d['h'], d['t'], d['i'], d['sk'], d['al'], d['persen']]
                for d in data]
        return table_response(['NISN', 'Nama', 'Kelas', 'Hadir', 'Terlambat', 'Izin', 'Sakit', 'Alpa',
                               '% Kehadiran'], rows, request.args.get('format', 'xlsx'),
                              f'rekap-{bulan}' + (f'-{kelas}' if kelas else ''), text_cols=(0,))

    @app.get('/admin/export')
    @login_required
    def export_day():
        tanggal = parse_date(request.args.get('tanggal'), today_str())
        kelas = request.args.get('kelas', '').strip()
        sql = '''SELECT s.nisn, s.nama, s.kelas, ? tanggal, COALESCE(a.jam_masuk,'-') jam,
                        COALESCE(a.status,'Belum absen') status, COALESCE(a.keterangan,'') ket
                 FROM siswa s LEFT JOIN absensi a ON a.siswa_id=s.id AND a.tanggal=? WHERE s.aktif=1'''
        params = [tanggal, tanggal]
        if kelas:
            sql += ' AND s.kelas=?'
            params.append(kelas)
        sql += ' ORDER BY s.kelas, s.nama COLLATE NOCASE'
        rows = [list(r) for r in get_db().execute(sql, params)]
        return table_response(['NISN', 'Nama', 'Kelas', 'Tanggal', 'Jam Masuk', 'Status', 'Keterangan'], rows,
                              request.args.get('format', 'xlsx'), f'absensi-{tanggal}', text_cols=(0,))

    @app.get('/admin/export-today')  # kompatibel dengan versi lama
    @login_required
    def export_today():
        return redirect(url_for('export_day', format='csv', tanggal=today_str()))

    # --------------------------------------------------------------- WA log
    @app.get('/admin/wa')
    @login_required
    def wa_log():
        status = request.args.get('status', '')
        db = get_db()
        sql = '''SELECT w.*, s.nama FROM wa_log w LEFT JOIN siswa s ON s.id=w.siswa_id'''
        params = []
        if status in ('pending', 'retry', 'sent', 'failed'):
            sql += ' WHERE w.status=?'
            params.append(status)
        rows = db.execute(sql + ' ORDER BY w.id DESC LIMIT 200', params).fetchall()
        summary = {r['status']: r['c'] for r in db.execute('SELECT status, COUNT(*) c FROM wa_log GROUP BY status')}
        return render_template('wa.html', rows=rows, status=status, summary=summary,
                               token_set=bool(cfg['FONNTE_TOKEN']))

    @app.post('/admin/wa/retry')
    @login_required
    def wa_retry():
        db = get_db()
        stamp = svc.now_str(cfg['TZ'])
        wid = request.form.get('id', type=int)
        if wid:
            n = db.execute("UPDATE wa_log SET status='pending', attempts=0, next_attempt_at=?, updated_at=? "
                           "WHERE id=? AND status IN ('failed','retry')", (stamp, stamp, wid)).rowcount
        else:
            n = db.execute("UPDATE wa_log SET status='pending', attempts=0, next_attempt_at=?, updated_at=? "
                           "WHERE status='failed'", (stamp, stamp)).rowcount
        db.commit()
        worker = app.extensions.get('worker')
        if worker:
            worker.wake.set()
        flash(f'{n} pesan dijadwalkan kirim ulang.', 'success')
        return redirect(url_for('wa_log'))

    # ----------------------------------------------------------- pengaturan
    @app.route('/admin/pengaturan', methods=['GET', 'POST'])
    @login_required
    def pengaturan():
        db = get_db()
        if request.method == 'POST':
            batas = request.form.get('batas_terlambat', '').strip()
            nama = request.form.get('nama_sekolah', '').strip()[:80]
            tpl = request.form.get('wa_template', '').strip()[:800]
            if not re.fullmatch(r'([01]\d|2[0-3]):[0-5]\d', batas):
                flash('Format jam batas harus HH:MM, contoh 07:15.', 'error')
            elif not nama or not tpl:
                flash('Nama sekolah dan template pesan wajib diisi.', 'error')
            else:
                dbm.set_setting(db, 'batas_terlambat', batas)
                dbm.set_setting(db, 'nama_sekolah', nama)
                dbm.set_setting(db, 'wa_template', tpl)
                dbm.set_setting(db, 'wa_aktif', '1' if request.form.get('wa_aktif') else '0')
                db.commit()
                flash('Pengaturan disimpan.', 'success')
            return redirect(url_for('pengaturan'))
        server_url = request.args.get('server_url', '').strip() or request.host_url.rstrip('/')
        setup_payload = 'ABSEN-SETUP:' + json.dumps({'u': server_url, 'k': cfg['SCAN_API_KEY']}, separators=(',', ':'))
        settings = {k: dbm.get_setting(db, k) for k in dbm.DEFAULT_SETTINGS}
        return render_template('pengaturan.html', s=settings, api_key=cfg['SCAN_API_KEY'], server_url=server_url,
                               setup_svg=qr_svg(setup_payload), is_local=urlparse(server_url).hostname in
                               ('127.0.0.1', 'localhost'), token_set=bool(cfg['FONNTE_TOKEN']),
                               tz=cfg['TIMEZONE'], legacy=cfg['ALLOW_LEGACY_NISN'])

    @app.post('/admin/password')
    @login_required
    def ganti_password():
        db = get_db()
        cur = request.form.get('lama', '')
        new = request.form.get('baru', '')
        if not check_login(cfg['ADMIN_USERNAME'], cur):
            flash('Password lama salah.', 'error')
        elif len(new) < 8 or new != request.form.get('ulang', ''):
            flash('Password baru minimal 8 karakter dan harus sama dengan konfirmasi.', 'error')
        else:
            dbm.set_setting(db, 'admin_password_hash', generate_password_hash(new))
            db.commit()
            flash('Password diganti. Password lama di .env tidak dipakai lagi.', 'success')
        return redirect(url_for('pengaturan'))

    @app.get('/admin/backup')
    @login_required
    def backup_download():
        stamp = svc.now_local(cfg['TZ']).strftime('%Y%m%d-%H%M%S')
        os.makedirs(cfg['BACKUP_DIR'], exist_ok=True)
        dest = os.path.join(cfg['BACKUP_DIR'], f'manual-{stamp}.db')
        svc.backup_db(cfg['DB_PATH'], dest)
        return send_file(dest, as_attachment=True, download_name=f'sekolah-{stamp}.db')

    return app


<<<<<<< HEAD
@app.get('/')
def home():
    conn = db()
    students = conn.execute('SELECT * FROM siswa ORDER BY nama COLLATE NOCASE').fetchall()
    
    # Hitung absensi hari ini
    today = now_local().strftime('%Y-%m-%d')
    attended_today = conn.execute(
        "SELECT DISTINCT siswa_id FROM absensi WHERE tanggal = ?", (today,)
    ).fetchall()
    attended_ids = {row['siswa_id'] for row in attended_today}
    
    students_data = []
    for s in students:
        s_dict = dict(s)
        s_dict['sudah_absen'] = s_dict['id'] in attended_ids
        students_data.append(s_dict)

    conn.close()
    return render_template('admin.html', students=students_data, attendance_today=len(attended_ids))


@app.get('/api/health')
def health():
    conn = db()
    total = conn.execute('SELECT COUNT(*) AS c FROM siswa WHERE aktif=1').fetchone()['c']
    today = now_local().strftime('%Y-%m-%d')
    hadir = conn.execute('SELECT COUNT(*) AS c FROM absensi WHERE tanggal=?', (today,)).fetchone()['c']
    conn.close()
    return jsonify({'ok': True, 'timezone': TIMEZONE, 'total_siswa_aktif': total, 'hadir_hari_ini': hadir})


@app.post('/scan')
def scan():
    data = request.get_json(silent=True) or {}
    nisn = str(data.get('nisn', '')).strip()
    device_id = str(data.get('device_id', '')).strip()[:100]

    if not nisn:
        return jsonify({'success': False, 'code': 'NISN_REQUIRED', 'message': 'NISN kosong.'}), 400

    conn = db()
    siswa = conn.execute(
        'SELECT * FROM siswa WHERE nisn=? AND aktif=1 LIMIT 1',
        (nisn,)
    ).fetchone()

    if not siswa:
        conn.close()
        return jsonify({'success': False, 'code': 'STUDENT_NOT_FOUND', 'message': 'NISN tidak terdaftar atau siswa tidak aktif.'}), 404

    current = now_local()
    tanggal = current.strftime('%Y-%m-%d')
    jam = current.strftime('%H:%M:%S')

    existing = conn.execute(
        '''SELECT id, jam_masuk, status FROM absensi WHERE siswa_id=? AND tanggal=? LIMIT 1''',
        (siswa['id'], tanggal)
    ).fetchone()

    if existing:
        conn.close()
        return jsonify({
            'success': True,
            'code': 'ALREADY_ATTENDED',
            'message': f"{siswa['nama']} sudah absen hari ini pada {existing['jam_masuk']}.",
            'student': dict(siswa),
            'attendance': dict(existing),
        })

    try:
        cur = conn.execute(
            '''INSERT INTO absensi (siswa_id, tanggal, jam_masuk, status, device_id, created_at)
               VALUES (?, ?, ?, 'Hadir', ?, ?)''',
            (siswa['id'], tanggal, jam, device_id, current.isoformat())
        )
        conn.commit()
        attendance_id = cur.lastrowid
    except sqlite3.IntegrityError:
        existing = conn.execute(
            'SELECT id, jam_masuk, status FROM absensi WHERE siswa_id=? AND tanggal=?',
            (siswa['id'], tanggal)
        ).fetchone()
        conn.close()
        return jsonify({
            'success': True,
            'code': 'ALREADY_ATTENDED',
            'message': f"{siswa['nama']} sudah absen hari ini.",
            'student': dict(siswa),
            'attendance': dict(existing) if existing else None,
        })

    wa_result = None
    whatsapp = (siswa['whatsapp_ortu'] or '').strip()
    if whatsapp:
        message = (
            f"PEMBERITAHUAN ABSENSI SEKOLAH\n\n"
            f"Nama: {siswa['nama']}\n"
            f"NISN: {siswa['nisn']}\n"
            f"Kelas: {siswa['kelas']}\n"
            f"Status: HADIR\n"
            f"Tanggal: {current.strftime('%d-%m-%Y')}\n"
            f"Jam masuk: {jam}\n\n"
            f"Pesan ini dikirim otomatis oleh sistem absensi sekolah."
        )
        wa_result = send_whatsapp(whatsapp, message)

    conn.close()
    return jsonify({
        'success': True,
        'code': 'ATTENDANCE_RECORDED',
        'message': f"Absensi {siswa['nama']} berhasil dicatat pada {jam}.",
        'student': {
            'nisn': siswa['nisn'],
            'nama': siswa['nama'],
            'kelas': siswa['kelas'],
        },
        'attendance': {
            'id': attendance_id,
            'tanggal': tanggal,
            'jam_masuk': jam,
            'status': 'Hadir',
        },
        'whatsapp': wa_result,
    }), 201


@app.get('/api/attendance/today')
def attendance_today():
    tanggal = now_local().strftime('%Y-%m-%d')
    conn = db()
    rows = conn.execute('''
        SELECT a.id, a.tanggal, a.jam_masuk, a.status, a.device_id,
               s.nisn, s.nama, s.kelas
        FROM absensi a
        JOIN siswa s ON s.id=a.siswa_id
        WHERE a.tanggal=?
        ORDER BY a.jam_masuk ASC
    ''', (tanggal,)).fetchall()
    conn.close()
    return jsonify({'date': tanggal, 'data': [dict(r) for r in rows]})


@app.post('/admin/siswa')
def add_student():
    nisn = request.form.get('nisn', '').strip()
    nama = request.form.get('nama', '').strip()
    kelas = request.form.get('kelas', '').strip()
    wa = request.form.get('whatsapp_ortu', '').strip()
    if not nisn or not nama or not kelas:
        flash('NISN, nama, dan kelas wajib diisi.', 'error')
        return redirect(url_for('home'))
    conn = db()
    try:
        conn.execute(
            'INSERT INTO siswa (nisn,nama,kelas,whatsapp_ortu,created_at) VALUES (?,?,?,?,?)',
            (nisn, nama, kelas, wa, now_local().isoformat())
        )
        conn.commit()
        flash('Siswa berhasil ditambahkan.', 'success')
    except sqlite3.IntegrityError:
        flash('NISN sudah terdaftar.', 'error')
    finally:
        conn.close()
    return redirect(url_for('home'))


@app.post('/admin/siswa/<int:siswa_id>/toggle')
def toggle_student(siswa_id):
    conn = db()
    row = conn.execute('SELECT aktif FROM siswa WHERE id=?', (siswa_id,)).fetchone()
    if row:
        conn.execute('UPDATE siswa SET aktif=? WHERE id=?', (0 if row['aktif'] else 1, siswa_id))
        conn.commit()
    conn.close()
    return redirect(url_for('home'))


@app.get('/admin/export-today')
def export_today():
    from csv import writer
    from io import StringIO
    tanggal = now_local().strftime('%Y-%m-%d')
    conn = db()
    rows = conn.execute('''
        SELECT s.nisn,s.nama,s.kelas,a.tanggal,a.jam_masuk,a.status
        FROM absensi a JOIN siswa s ON s.id=a.siswa_id WHERE a.tanggal=?
        ORDER BY a.jam_masuk
    ''', (tanggal,)).fetchall()
    conn.close()
    out = StringIO()
    w = writer(out)
    w.writerow(['NISN', 'Nama', 'Kelas', 'Tanggal', 'Jam Masuk', 'Status'])
    for r in rows:
        w.writerow([r['nisn'], r['nama'], r['kelas'], r['tanggal'], r['jam_masuk'], r['status']])
    return Response(out.getvalue(), mimetype='text/csv; charset=utf-8', headers={
        'Content-Disposition': f'attachment; filename="absensi-{tanggal}.csv"'
    })


@app.get('/api/qr/<nisn>')
def qr_payload(nisn):
    conn = db()
    row = conn.execute('SELECT nisn,nama,kelas FROM siswa WHERE nisn=? AND aktif=1', (nisn,)).fetchone()
    conn.close()
    if not row:
        return jsonify({'success': False, 'message': 'Siswa tidak ditemukan.'}), 404
    return jsonify({'success': True, 'qr_value': row['nisn'], 'student': dict(row)})
=======
def main():
    cfgmod.ensure_secrets()
    app = create_app()
    cfg = app.extensions['absensi_cfg']
    worker = svc.BackgroundWorker(cfg)
    app.extensions['worker'] = worker
    worker.start()
    ip = guess_lan_ip()
    log.info('Server siap. Admin: http://%s:%s  (HP: gunakan alamat ini di aplikasi)', ip, cfg['PORT'])
    try:
        from waitress import serve
    except ImportError:  # pragma: no cover
        log.warning('waitress tidak terpasang; memakai server pengembangan Flask.')
        app.run(host=cfg['HOST'], port=cfg['PORT'], debug=False)
    else:
        serve(app, host=cfg['HOST'], port=cfg['PORT'], threads=8)
>>>>>>> 200af4a (Absensi Siswa v2)


# ==========================================
# ROUTE PERBAIKAN: EDIT DATA SISWA
# ==========================================
@app.route('/admin/siswa/<int:id>/edit', methods=['POST'])
def edit_siswa(id):
    nisn = request.form.get('nisn', '').strip()
    nama = request.form.get('nama', '').strip()
    kelas = request.form.get('kelas', '').strip()
    wa = request.form.get('whatsapp_ortu', '').strip()
    
    conn = db()
    try:
        conn.execute(
            "UPDATE siswa SET nisn=?, nama=?, kelas=?, whatsapp_ortu=? WHERE id=?", 
            (nisn, nama, kelas, wa, id)
        )
        conn.commit()
        flash('Data siswa berhasil diperbarui!', 'success')
    except sqlite3.IntegrityError:
        flash('Gagal memperbarui: NISN sudah digunakan oleh siswa lain.', 'error')
    finally:
        conn.close()
    return redirect(url_for('home'))


# ==========================================
# ROUTE PERBAIKAN: HAPUS SISWA
# ==========================================
@app.route('/admin/siswa/<int:id>/delete', methods=['POST'])
def delete_siswa(id):
    conn = db()
    conn.execute("DELETE FROM siswa WHERE id=?", (id,))
    conn.commit()
    conn.close()
    flash('Siswa berhasil dihapus!', 'success')
    return redirect(url_for('home'))


if __name__ == '__main__':
<<<<<<< HEAD
    init_db()
    host = os.getenv('FLASK_HOST', '0.0.0.0')
    port = int(os.getenv('FLASK_PORT', '5000'))
    app.run(host=host, port=port, debug=False)
=======
    main()
>>>>>>> 200af4a (Absensi Siswa v2)
