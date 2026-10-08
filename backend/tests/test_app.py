import io
import re
import sqlite3

import pytest

import database as dbm
import services as svc
from app import create_app

API_KEY = 'kunci-uji-123'
PASSWORD = 'password-uji-123'


@pytest.fixture()
def app(tmp_path):
    return create_app({
        'DB_PATH': str(tmp_path / 't.db'), 'BACKUP_DIR': str(tmp_path / 'bk'),
        'UPLOAD_DIR': str(tmp_path / 'up'), 'LOG_DIR': str(tmp_path / 'lg'), 'LOG_TO_FILE': False,
        'FLASK_SECRET_KEY': 'rahasia-uji', 'QR_SECRET': 'qr-rahasia-uji', 'SCAN_API_KEY': API_KEY,
        'ADMIN_USERNAME': 'admin', 'ADMIN_PASSWORD': PASSWORD, 'FONNTE_TOKEN': 'tok',
        'TIMEZONE': 'Asia/Jakarta',
    })


@pytest.fixture()
def cfg(app):
    return app.extensions['absensi_cfg']


@pytest.fixture()
def client(app):
    return app.test_client()


def csrf_of(client, path='/login'):
    html = client.get(path).get_data(as_text=True)
    m = re.search(r'name="csrf_token" value="([0-9a-f]+)"', html)
    assert m, 'token csrf tidak ditemukan'
    return m.group(1)


@pytest.fixture()
def admin(client):
    tok = csrf_of(client)
    r = client.post('/login', data={'username': 'admin', 'password': PASSWORD, 'csrf_token': tok})
    assert r.status_code == 302
    return client


def post(client, path, data=None, page='/'):
    d = dict(data or {})
    d['csrf_token'] = csrf_of(client, page)
    return client.post(path, data=d)


def add_student(admin, nisn='0012345678', nama='Budi', kelas='X IPA 1', wa='081234567890'):
    return post(admin, '/admin/siswa', {'nisn': nisn, 'nama': nama, 'kelas': kelas, 'whatsapp_ortu': wa},
                '/admin/siswa')


def qr_for(cfg, nisn='0012345678', version=1):
    return svc.make_qr_payload(cfg['QR_SECRET'], nisn, version)


def scan(client, qr, key=API_KEY, **extra):
    headers = {'X-API-Key': key} if key is not None else {}
    return client.post('/scan', json={'qr': qr, 'device_id': 'HP1', **extra}, headers=headers)


# ------------------------------------------------------------------ auth
def test_admin_pages_require_login(client):
    for path in ['/', '/admin/siswa', '/admin/rekap', '/admin/kartu', '/admin/wa', '/admin/pengaturan',
                 '/admin/backup', '/admin/export']:
        r = client.get(path)
        assert r.status_code == 302 and '/login' in r.headers['Location'], path
    assert client.get('/api/attendance/today').status_code == 401
    assert client.get('/api/qr/0012345678').status_code == 401


def test_login_wrong_and_right(client):
    tok = csrf_of(client)
    r = client.post('/login', data={'username': 'admin', 'password': 'salah', 'csrf_token': tok})
    assert r.status_code == 200 and 'salah' in r.get_data(as_text=True)
    tok = csrf_of(client)
    r = client.post('/login', data={'username': 'admin', 'password': PASSWORD, 'csrf_token': tok})
    assert r.status_code == 302
    assert client.get('/').status_code == 200


def test_login_rate_limit(client):
    for _ in range(5):
        tok = csrf_of(client)
        client.post('/login', data={'username': 'admin', 'password': 'x', 'csrf_token': tok})
    tok = csrf_of(client)
    r = client.post('/login', data={'username': 'admin', 'password': PASSWORD, 'csrf_token': tok})
    assert r.status_code == 429


def test_csrf_required_on_post(admin):
    r = admin.post('/admin/siswa', data={'nisn': '0012345678', 'nama': 'A', 'kelas': 'X'})
    assert r.status_code == 400


def test_open_redirect_blocked(client):
    tok = csrf_of(client)
    r = client.post('/login?next=//evil.com', data={'username': 'admin', 'password': PASSWORD, 'csrf_token': tok})
    assert r.headers['Location'].endswith('/')
    assert 'evil.com' not in r.headers['Location']


# ------------------------------------------------------------------ scan
def test_scan_requires_api_key(client, cfg):
    assert scan(client, qr_for(cfg), key=None).status_code == 401
    assert scan(client, qr_for(cfg), key='salah').status_code == 401


def test_scan_flow_and_duplicate(admin, cfg):
    add_student(admin)
    r = scan(admin, qr_for(cfg))
    assert r.status_code == 201
    j = r.get_json()
    assert j['code'] == 'ATTENDANCE_RECORDED' and j['student']['nama'] == 'Budi'
    assert j['attendance']['status'] in ('Hadir', 'Terlambat')
    assert j['whatsapp']['queued'] is True
    r2 = scan(admin, qr_for(cfg))
    assert r2.status_code == 200 and r2.get_json()['code'] == 'ALREADY_ATTENDED'
    conn = dbm.connect(cfg['DB_PATH'])
    assert conn.execute('SELECT COUNT(*) c FROM absensi').fetchone()['c'] == 1
    assert conn.execute('SELECT COUNT(*) c FROM wa_log').fetchone()['c'] == 1


def test_scan_rejects_forged_and_legacy(admin, cfg):
    add_student(admin)
    assert scan(admin, 'ABSEN:0012345678:000000000000').get_json()['code'] == 'QR_INVALID'
    assert scan(admin, '0012345678').get_json()['code'] == 'QR_LEGACY'
    assert scan(admin, 'NISN:0012345678').status_code == 422
    assert scan(admin, 'hello').get_json()['code'] == 'QR_INVALID'
    assert scan(admin, 'ABSEN:0099999999:' + 'a' * 12).get_json()['code'] == 'STUDENT_NOT_FOUND'
    cfg['ALLOW_LEGACY_NISN'] = True
    assert scan(admin, '0012345678').status_code == 201


def test_old_json_field_nisn_is_legacy(admin, cfg):
    add_student(admin)
    r = admin.post('/scan', json={'nisn': '0012345678'}, headers={'X-API-Key': API_KEY})
    assert r.get_json()['code'] == 'QR_LEGACY'


def test_reset_kartu_invalidates_old_qr(admin, cfg):
    add_student(admin)
    old = qr_for(cfg)
    sid = dbm.connect(cfg['DB_PATH']).execute('SELECT id FROM siswa').fetchone()['id']
    post(admin, f'/admin/siswa/{sid}/reset-kartu', page='/admin/siswa')
    assert scan(admin, old).get_json()['code'] == 'QR_INVALID'
    assert scan(admin, qr_for(cfg, version=2)).status_code == 201


def test_inactive_student_rejected(admin, cfg):
    add_student(admin)
    sid = dbm.connect(cfg['DB_PATH']).execute('SELECT id FROM siswa').fetchone()['id']
    post(admin, f'/admin/siswa/{sid}/toggle', page='/admin/siswa')
    assert scan(admin, qr_for(cfg)).status_code == 404


def test_late_status(admin, cfg):
    add_student(admin)
    conn = dbm.connect(cfg['DB_PATH'])
    dbm.set_setting(conn, 'batas_terlambat', '00:00')
    conn.commit()
    assert scan(admin, qr_for(cfg)).get_json()['attendance']['status'] == 'Terlambat'


def test_scan_after_alpa_converts_to_present(admin, cfg):
    add_student(admin)
    conn = dbm.connect(cfg['DB_PATH'])
    n = svc.mark_alpa(conn, cfg, svc.now_local(cfg['TZ']).strftime('%Y-%m-%d'))
    conn.commit()
    assert n == 1
    r = scan(admin, qr_for(cfg))
    assert r.status_code == 201
    row = conn.execute('SELECT status FROM absensi').fetchone()
    assert row['status'] in ('Hadir', 'Terlambat')
    assert conn.execute('SELECT COUNT(*) c FROM absensi').fetchone()['c'] == 1


def test_scan_does_not_require_csrf_or_session(app, cfg):
    c = app.test_client()
    conn = dbm.connect(cfg['DB_PATH'])
    conn.execute("INSERT INTO siswa (nisn,nama,kelas,created_at) VALUES ('0012345678','B','X','x')")
    conn.commit()
    assert scan(c, qr_for(cfg)).status_code == 201


# -------------------------------------------------------------------- WA
def test_wa_queue_success_and_template(admin, cfg):
    add_student(admin)
    conn = dbm.connect(cfg['DB_PATH'])
    dbm.set_setting(conn, 'batas_terlambat', '23:59')  # deterministik: selalu 'Hadir'
    conn.commit()
    scan(admin, qr_for(cfg))
    sent = []

    def ok_sender(token, target, message, cc):
        sent.append((target, message))
        return True, 'ok'
    assert svc.process_wa_queue(cfg, sender=ok_sender) == 1
    assert sent[0][0] == '6281234567890' and 'Budi' in sent[0][1] and 'HADIR' in sent[0][1].upper()
    conn = dbm.connect(cfg['DB_PATH'])
    assert conn.execute('SELECT status FROM wa_log').fetchone()['status'] == 'sent'
    assert svc.process_wa_queue(cfg, sender=ok_sender) == 0


def test_wa_retry_then_failed(admin, cfg):
    add_student(admin)
    scan(admin, qr_for(cfg))
    conn = dbm.connect(cfg['DB_PATH'])
    fail = lambda *a: (False, 'boom')  # noqa: E731
    for expected in ('retry', 'retry', 'failed'):
        conn.execute("UPDATE wa_log SET next_attempt_at='2000-01-01 00:00:00' WHERE status IN ('pending','retry')")
        conn.commit()
        svc.process_wa_queue(cfg, sender=fail)
        assert conn.execute('SELECT status FROM wa_log').fetchone()['status'] == expected
    r = post(admin, '/admin/wa/retry', page='/admin/wa')
    assert r.status_code == 302
    assert conn.execute('SELECT status FROM wa_log').fetchone()['status'] == 'pending'


def test_fonnte_http200_status_false_is_failure(monkeypatch):
    class R:
        ok = True
        text = ''
        def json(self): return {'status': False, 'reason': 'invalid token'}
    monkeypatch.setattr(svc.requests, 'post', lambda *a, **k: R())
    ok, detail = svc.send_fonnte('t', '62812', 'hi')
    assert ok is False and 'invalid token' in detail


def test_wa_without_token_marks_failed(admin, cfg):
    add_student(admin)
    scan(admin, qr_for(cfg))
    cfg['FONNTE_TOKEN'] = ''
    svc.process_wa_queue(cfg)
    row = dbm.connect(cfg['DB_PATH']).execute('SELECT * FROM wa_log').fetchone()
    assert row['status'] == 'failed' and 'FONNTE_TOKEN' in row['last_response']


def test_wa_disabled_or_no_number_not_queued(admin, cfg):
    add_student(admin, wa='')
    assert scan(admin, qr_for(cfg)).get_json()['whatsapp']['queued'] is False


def test_normalizers():
    assert svc.normalize_wa('0812-3456-7890') == '6281234567890'
    assert svc.normalize_wa('+62 812 3456 7890') == '6281234567890'
    assert svc.normalize_wa('81234567890') == '6281234567890'
    assert svc.normalize_wa('') == ''
    assert svc.normalize_wa('12345') is None
    assert svc.normalize_nisn('12345678') == '0012345678'
    assert svc.normalize_nisn('12345678901') is None
    assert svc.normalize_nisn('abc') is None
    assert svc.render_message('Hai {nama} {x}', {'nama': 'A'}) == 'Hai A {x}'


# ------------------------------------------------------- admin: dashboard
def test_manual_status_and_dashboard(admin, cfg):
    add_student(admin)
    sid = dbm.connect(cfg['DB_PATH']).execute('SELECT id FROM siswa').fetchone()['id']
    today = svc.now_local(cfg['TZ']).strftime('%Y-%m-%d')
    r = post(admin, '/admin/absensi/set', {'siswa_id': sid, 'tanggal': today, 'status': 'Sakit',
                                          'keterangan': 'demam'})
    assert r.status_code == 302
    html = admin.get('/').get_data(as_text=True)
    assert 'Sakit' in html and 'demam' in html
    # tanggal masa depan ditolak
    post(admin, '/admin/absensi/set', {'siswa_id': sid, 'tanggal': '2999-01-01', 'status': 'Izin'})
    assert dbm.connect(cfg['DB_PATH']).execute('SELECT COUNT(*) c FROM absensi').fetchone()['c'] == 1


# --------------------------------------------------------- admin: siswa
def test_student_crud_search_pagination(admin, cfg):
    for i in range(30):
        add_student(admin, nisn=f'{i + 1:010d}', nama=f'Siswa {i:02d}', kelas='X A' if i < 20 else 'X B', wa='')
    assert 'Halaman 1 / 2' in admin.get('/admin/siswa').get_data(as_text=True)
    assert 'Siswa 05' in admin.get('/admin/siswa?q=Siswa 05').get_data(as_text=True)
    assert 'Siswa 25' not in admin.get('/admin/siswa?kelas=X A').get_data(as_text=True)
    r = add_student(admin, nisn='0000000001')  # duplikat
    assert 'sudah terdaftar' in admin.get('/admin/siswa').get_data(as_text=True)
    sid = dbm.connect(cfg['DB_PATH']).execute("SELECT id FROM siswa WHERE nisn='0000000001'").fetchone()['id']
    r = post(admin, f'/admin/siswa/{sid}/edit', {'nisn': '0000000001', 'nama': 'Baru', 'kelas': 'XI A',
                                                'whatsapp_ortu': '0812345678901'}, f'/admin/siswa/{sid}/edit')
    assert r.status_code == 302
    row = dbm.connect(cfg['DB_PATH']).execute('SELECT * FROM siswa WHERE id=?', (sid,)).fetchone()
    assert row['nama'] == 'Baru' and row['whatsapp_ortu'] == '62812345678901'
    post(admin, f'/admin/siswa/{sid}/delete', page='/admin/siswa')
    assert dbm.connect(cfg['DB_PATH']).execute('SELECT COUNT(*) c FROM siswa').fetchone()['c'] == 29


def test_validation_rejects_bad_input(admin, cfg):
    add_student(admin, nisn='abc')
    add_student(admin, wa='123')
    assert dbm.connect(cfg['DB_PATH']).execute('SELECT COUNT(*) c FROM siswa').fetchone()['c'] == 0


def test_kelas_aksi(admin, cfg):
    add_student(admin, nisn='0000000001', kelas='XII A')
    add_student(admin, nisn='0000000002', kelas='XII A')
    post(admin, '/admin/siswa/kelas-aksi', {'aksi': 'nonaktifkan', 'dari': 'XII A'}, '/admin/siswa')
    assert dbm.connect(cfg['DB_PATH']).execute('SELECT COUNT(*) c FROM siswa WHERE aktif=1').fetchone()['c'] == 0
    post(admin, '/admin/siswa/kelas-aksi', {'aksi': 'pindah', 'dari': 'XII A', 'ke': 'Alumni'}, '/admin/siswa')
    assert dbm.connect(cfg['DB_PATH']).execute("SELECT COUNT(*) c FROM siswa WHERE kelas='Alumni'").fetchone()['c'] == 2


def test_photo_upload_and_serve(admin, cfg):
    from PIL import Image
    add_student(admin)
    sid = dbm.connect(cfg['DB_PATH']).execute('SELECT id FROM siswa').fetchone()['id']
    buf = io.BytesIO()
    Image.new('RGB', (1000, 800), 'red').save(buf, 'PNG')
    buf.seek(0)
    tok = csrf_of(admin, f'/admin/siswa/{sid}/edit')
    r = admin.post(f'/admin/siswa/{sid}/edit', data={
        'csrf_token': tok, 'nisn': '0012345678', 'nama': 'Budi', 'kelas': 'X IPA 1',
        'whatsapp_ortu': '', 'foto': (buf, 'a.png')}, content_type='multipart/form-data')
    assert r.status_code == 302
    r = admin.get(f'/foto/{sid}')
    assert r.status_code == 200 and r.mimetype == 'image/jpeg'
    assert Image.open(io.BytesIO(r.data)).size[0] <= 400
    j = scan(admin, qr_for(cfg)).get_json()
    assert j['student']['foto_url'].startswith(f'/foto/{sid}')
    # perangkat dengan API key boleh ambil foto, tanpa kunci tidak
    anon = admin.application.test_client()
    assert anon.get(f'/foto/{sid}').status_code == 404
    assert anon.get(f'/foto/{sid}', headers={'X-API-Key': API_KEY}).status_code == 200
    # bukan gambar ditolak
    tok = csrf_of(admin, f'/admin/siswa/{sid}/edit')
    admin.post(f'/admin/siswa/{sid}/edit', data={
        'csrf_token': tok, 'nisn': '0012345678', 'nama': 'Budi', 'kelas': 'X IPA 1', 'whatsapp_ortu': '',
        'foto': (io.BytesIO(b'bukan gambar'), 'x.jpg')}, content_type='multipart/form-data')
    assert 'bukan gambar yang valid' in admin.get('/admin/siswa').get_data(as_text=True)


# --------------------------------------------------------------- import
def test_import_csv_and_xlsx(admin, cfg):
    csv_data = ('NISN;Nama;Kelas;WhatsApp\n'
                '12345678;Ani;X A;0812345678901\n'
                '0000000002;Budi;X A;\n'
                'xx;Salah;X A;\n'
                '0000000003;Cici;;\n').encode()
    tok = csrf_of(admin, '/admin/import')
    r = admin.post('/admin/import', data={'csrf_token': tok, 'perbarui': '1',
                                         'berkas': (io.BytesIO(csv_data), 'siswa.csv')},
                   content_type='multipart/form-data')
    html = r.get_data(as_text=True)
    assert 'Ditambahkan: <b>2</b>' in html and 'Dilewati: <b>2</b>' in html
    conn = dbm.connect(cfg['DB_PATH'])
    assert conn.execute("SELECT nama FROM siswa WHERE nisn='0012345678'").fetchone()['nama'] == 'Ani'

    # xlsx dari template (angka tanpa nol di depan)
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(['NISN', 'Nama', 'Kelas', 'WhatsApp Ortu'])
    ws.append([12345678, 'Ani Baru', 'X B', 81234567890])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    tok = csrf_of(admin, '/admin/import')
    admin.post('/admin/import', data={'csrf_token': tok, 'perbarui': '1', 'berkas': (buf, 's.xlsx')},
               content_type='multipart/form-data')
    row = conn.execute("SELECT * FROM siswa WHERE nisn='0012345678'").fetchone()
    assert row['nama'] == 'Ani Baru' and row['whatsapp_ortu'] == '6281234567890'


def test_import_template_download(admin):
    r = admin.get('/admin/import/template')
    assert r.status_code == 200 and r.data[:2] == b'PK'


# ---------------------------------------------------------- kartu & rekap
def test_kartu_page_has_signed_qr(admin, cfg):
    add_student(admin)
    html = admin.get('/admin/kartu').get_data(as_text=True)
    assert '<svg' in html and 'Budi' in html and '0012345678' in html


def test_rekap_and_exports(admin, cfg):
    add_student(admin, nisn='0000000001', nama='A')
    add_student(admin, nisn='0000000002', nama='B')
    conn = dbm.connect(cfg['DB_PATH'])
    ids = {r['nisn']: r['id'] for r in conn.execute('SELECT id, nisn FROM siswa')}
    for d, st in [('2026-03-02', 'Hadir'), ('2026-03-03', 'Terlambat'), ('2026-03-04', 'Hadir')]:
        svc.set_attendance(conn, cfg, ids['0000000001'], d, st)
    svc.set_attendance(conn, cfg, ids['0000000002'], '2026-03-02', 'Hadir')
    svc.set_attendance(conn, cfg, ids['0000000002'], '2026-03-03', 'Sakit')
    conn.commit()
    html = admin.get('/admin/rekap?bulan=2026-03').get_data(as_text=True)
    assert 'Hari efektif bulan ini: <b>3</b>' in html and '100.0%' in html and '33.3%' in html
    x = admin.get('/admin/rekap/export?bulan=2026-03')
    assert x.data[:2] == b'PK'
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(x.data)).active
    assert ws['A2'].value == '0000000001' and ws['A2'].number_format == '@'
    c = admin.get('/admin/rekap/export?bulan=2026-03&format=csv')
    assert c.data.startswith(b'\xef\xbb\xbf') and b'Terlambat' in c.data
    d = admin.get('/admin/export?tanggal=2026-03-02&format=csv')
    assert b'Hadir' in d.data
    assert admin.get('/admin/export-today').status_code == 302


def test_csv_formula_injection_neutralised(admin, cfg):
    add_student(admin, nama='=HYPERLINK("http://x")')
    c = admin.get('/admin/export?format=csv')
    assert b"'=HYPERLINK" in c.data


# ----------------------------------------------------- pengaturan & backup
def test_settings_password_backup(admin, cfg):
    r = post(admin, '/admin/pengaturan', {'nama_sekolah': 'SMA Contoh', 'batas_terlambat': '07:30',
                                         'wa_template': 'Hai {nama}', 'wa_aktif': '1'}, '/admin/pengaturan')
    assert r.status_code == 302
    html = admin.get('/admin/pengaturan').get_data(as_text=True)
    assert 'SMA Contoh' in html and '07:30' in html and API_KEY in html
    post(admin, '/admin/pengaturan', {'nama_sekolah': 'X', 'batas_terlambat': '7:5', 'wa_template': 'x'},
         '/admin/pengaturan')
    assert dbm.get_setting(dbm.connect(cfg['DB_PATH']), 'batas_terlambat') == '07:30'

    post(admin, '/admin/password', {'lama': PASSWORD, 'baru': 'baru-12345', 'ulang': 'baru-12345'}, '/admin/pengaturan')
    other = admin.application.test_client()
    tok = csrf_of(other)
    assert other.post('/login', data={'username': 'admin', 'password': PASSWORD, 'csrf_token': tok}).status_code == 200
    tok = csrf_of(other)
    assert other.post('/login', data={'username': 'admin', 'password': 'baru-12345', 'csrf_token': tok}).status_code == 302

    r = admin.get('/admin/backup')
    assert r.status_code == 200 and r.data[:15] == b'SQLite format 3'


def test_daily_backup_and_prune(cfg, app):
    assert svc.maybe_daily_backup(cfg) is not None
    assert svc.maybe_daily_backup(cfg) is None


# ------------------------------------------------------------- migrasi
def test_migration_from_v1_database(tmp_path):
    path = str(tmp_path / 'old.db')
    conn = sqlite3.connect(path)
    conn.executescript('''
    CREATE TABLE siswa (id INTEGER PRIMARY KEY AUTOINCREMENT, nisn TEXT NOT NULL UNIQUE, nama TEXT NOT NULL,
      kelas TEXT NOT NULL, whatsapp_ortu TEXT, aktif INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL);
    CREATE TABLE absensi (id INTEGER PRIMARY KEY AUTOINCREMENT, siswa_id INTEGER NOT NULL, tanggal TEXT NOT NULL,
      jam_masuk TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'Hadir', device_id TEXT, created_at TEXT NOT NULL,
      FOREIGN KEY(siswa_id) REFERENCES siswa(id) ON DELETE CASCADE, UNIQUE(siswa_id, tanggal));
    INSERT INTO siswa (nisn,nama,kelas,whatsapp_ortu,created_at) VALUES ('0012345678','Lama','X','0812','2025-01-01');
    INSERT INTO absensi (siswa_id,tanggal,jam_masuk,created_at) VALUES (1,'2025-01-02','07:00:00','x');
    ''')
    conn.commit()
    conn.close()
    dbm.init_db(path)
    dbm.init_db(path)  # idempoten
    c = dbm.connect(path)
    assert c.execute('SELECT nama, qr_version FROM siswa').fetchone()['qr_version'] == 1
    assert c.execute('SELECT keterangan FROM absensi').fetchone()['keterangan'] is None
    assert c.execute('SELECT COUNT(*) c FROM pengaturan').fetchone()['c'] >= 4


def test_timezone_errors_are_explicit():
    from config import get_tz
    assert get_tz('Asia/Jakarta') is not None
    with pytest.raises(RuntimeError):
        get_tz('Mars/Olympus')


def test_ensure_secrets(tmp_path, monkeypatch):
    import config
    for k in ('FLASK_SECRET_KEY', 'QR_SECRET', 'SCAN_API_KEY', 'ADMIN_PASSWORD'):
        monkeypatch.delenv(k, raising=False)
    env = tmp_path / '.env'
    env.write_text('FLASK_SECRET_KEY=ganti-dengan-random-string\nFONNTE_TOKEN=abc\nSCAN_API_KEY=sudah-ada-123456\n')
    out = []
    got = config.ensure_secrets(str(env), announce=out.append)
    assert set(got) == {'FLASK_SECRET_KEY', 'QR_SECRET', 'ADMIN_PASSWORD'}
    text = env.read_text()
    assert 'FONNTE_TOKEN=abc' in text and 'SCAN_API_KEY=sudah-ada-123456' in text
    assert 'ganti-dengan' not in text
    assert config.ensure_secrets(str(env), announce=out.append) == {}


def test_health_and_ping(client):
    assert client.get('/api/health').get_json() == {'ok': True}
    assert client.get('/api/device/ping').status_code == 401
    j = client.get('/api/device/ping', headers={'X-API-Key': API_KEY}).get_json()
    assert j['ok'] is True and 'sekolah' in j


def test_all_admin_pages_render(admin, cfg):
    add_student(admin)
    scan(admin, qr_for(cfg))
    sid = dbm.connect(cfg['DB_PATH']).execute('SELECT id FROM siswa').fetchone()['id']
    for path in ['/', '/admin/siswa', f'/admin/siswa/{sid}/edit', '/admin/rekap', '/admin/kartu',
                 f'/admin/kartu?id={sid}', '/admin/import', '/admin/wa', '/admin/wa?status=pending',
                 '/admin/pengaturan', '/admin/pengaturan?server_url=http://192.168.1.5:5000',
                 '/api/attendance/today', f'/api/qr/0012345678', '/nonexistent']:
        r = admin.get(path)
        assert r.status_code in (200, 404), (path, r.status_code)
        if path != '/nonexistent':
            assert r.status_code == 200, path
